# ============================================================
# UPLOAD + SANDBOX IMPORT
# ============================================================
# Uploaded databases, SQL scripts and Excel files are stored and
# processed in a sandbox directory. Imported SQL never runs
# against store.db and uploaded SQL scripts only execute against
# a brand-new throwaway database.
#
# PERSISTENCE
# Every import is written through app.storage rather than straight to
# the filesystem. That indirection is what lets a durable backend be
# configured for production, where the local filesystem is read-only
# and /tmp does not survive an instance. This module owns the import
# SEMANTICS (validation, conversion, schema inspection, the namespace
# rules); app.storage owns WHERE the bytes and their records live.
#
# An import is now two things that must both be present to resolve:
#
#   * an object, stored under a key derived from the caller's namespace
#     and the generated id (for example "databases/<id>_<name>" or
#     "visitors/<id>/sandboxes/<id>_sandbox.db");
#   * a metadata record naming its owner (scope), kind, tables and the
#     object key.
#
# The record is authoritative. Lookups go through it, so a file whose
# name matches is NOT enough to resolve an import: it must have been
# recorded by an import in the same scope. Namespace isolation is thus
# enforced by the scope filter, not only by directory layout, which
# survives a backend that has no directories at all.
# ============================================================

import csv
import io
import os
import re
import sqlite3
import tempfile
import time
import uuid

import app.config as config
from app import storage
from app.schema import (
    inspect_database,
    schema_to_text,
)
from app.validate import ValidationError


class ImportNotFound(ValidationError):
    """A supplied import_id is well formed but matches no stored import.

    Subclasses ValidationError because it is still a deliberate,
    user-facing message. Call sites that want a 404 rather than a 400
    can catch this type specifically; everything that only knows about
    ValidationError keeps working unchanged.
    """


_TABLE_NAME_RE = re.compile(r"^[A-Za-z_][A-Za-z0-9_]*$")
_SAFE_ID_RE = re.compile(r"[^A-Za-z0-9_.-]")

# A stored import id is uuid4().hex[:12]: lowercase hex. The pattern is
# deliberately narrow so a value that could ever be a path fragment
# (a slash, a dot, a backslash, a drive colon) is rejected before any
# storage call is made.
_IMPORT_ID_RE = re.compile(r"^[A-Za-z0-9]{1,64}$")

# An anonymous visitor id is new_visitor_id(): 32 hex characters. The
# same reasoning as the import-id rule applies one level up: the id
# names a namespace, so anything that could escape it is refused before
# the key is built.
_VISITOR_ID_RE = re.compile(r"^[A-Za-z0-9]{16,64}$")

# The whole scope a visitor's imports live in. It must be exactly the
# visitor subtree, so a caller cannot aim an import at the shared
# namespace or somewhere above the root by way of the root argument.
_VISITOR_SCOPE_RE = re.compile(r"^visitors/[A-Za-z0-9]{16,64}$")

_STATEMENT_START_RE = re.compile(
    r"^(CREATE\s+(TABLE|UNIQUE\s+INDEX|INDEX|VIEW)|INSERT\s+INTO)\b",
    re.IGNORECASE,
)


def _upload_root():
    return os.path.abspath(config.UPLOAD_DIR)


# Every visitor's files live under their own subdirectory, so the import
# table of one browser can never be reached by another. The id is
# validated before it is joined to any path.
def _visitor_root(visitor_id):
    if not visitor_id or not _VISITOR_ID_RE.match(visitor_id):
        raise ValidationError("Invalid visitor id.")
    return os.path.join(_upload_root(), "visitors", visitor_id)


# ============================================================
# STORAGE BINDING
# ============================================================
# The stores are fetched per call (app.storage.get_stores), never
# cached, so a misconfigured backend fails closed on first use and a
# test that swaps the stores is not masked by a stale reference.

def _stores():
    return storage.get_stores()


def _object_store():
    return _stores().object_store


def _metadata_store():
    return _stores().metadata_store


def _key(prefix, sub, name):
    """Build an object key from a namespace prefix and a subdirectory."""

    parts = [part for part in (prefix, sub, name) if part]
    return "/".join(parts)


def _scope_for(root):
    """Turn a caller's root into a (scope, prefix, base) triple.

    ``scope`` is the durable owner label ("shared" or a visitor's
    subtree), ``prefix`` is the matching object-key prefix ("" for the
    shared namespace) and ``base`` is the absolute local root. Only the
    shared namespace and a single visitor's subtree are accepted, so an
    unexpected root cannot smuggle an import into another namespace.
    """

    base = os.path.abspath(root or _upload_root())
    root_abs = _upload_root()

    try:
        relative = os.path.relpath(base, root_abs)
    except ValueError:
        raise ValidationError("Invalid import namespace.")

    relative = relative.replace(os.sep, "/")

    if relative in (".", ""):
        return "shared", "", base

    if relative == ".." or relative.startswith("../"):
        raise ValidationError("Invalid import namespace.")

    if not _VISITOR_SCOPE_RE.match(relative):
        raise ValidationError("Invalid import namespace.")

    return relative, relative, base


def _save_import(
    import_id,
    scope,
    prefix,
    kind,
    filename,
    stored_name,
    sub,
    size,
    inspected,
    extra,
):
    """Write the object's metadata record; the object is already stored."""

    record = {
        "id": import_id,
        "scope": scope,
        "kind": kind,
        "filename": filename,
        "stored_name": stored_name,
        "object_key": _key(prefix, sub, stored_name),
        "path": f"{sub}/{stored_name}",
        "size": size,
        "created_at": time.time(),
        "tables": inspected.get("tables", []),
        "table_count": inspected.get("table_count", 0),
        "column_count": inspected.get("column_count", 0),
        "schema_text": inspected.get("schema_text", ""),
    }

    record.update(extra)

    _metadata_store().save(scope, record)

    return record


def _new_id():
    return uuid.uuid4().hex[:12]


def _sanitize_filename(filename):
    safe = _SAFE_ID_RE.sub("_", os.path.basename(filename or "upload"))
    return safe or "upload"


def _ensure_upload_size(content):
    maximum = config.MAX_UPLOAD_MB * 1024 * 1024
    if len(content) > maximum:
        raise ValidationError(
            f"Upload is too large. Maximum is {config.MAX_UPLOAD_MB} MB."
        )
    return content


def inspect_upload(path):
    """Inspect an importable database file (db or converted sandbox)."""

    schema = inspect_database(path)
    return {
        "schema": schema,
        "schema_text": schema_to_text(schema),
        "tables": [t["name"] for t in schema["tables"]],
        "table_count": schema["table_count"],
        "column_count": schema["column_count"],
    }


def _read_object_bytes(path):
    with open(path, "rb") as handle:
        return handle.read()


# ============================================================
# SQLITE DATABASE FILES
# ============================================================

def import_database(filename, content, root=None):
    """Copy an uploaded .db into the namespace and inspect it."""

    _ensure_upload_size(content)

    scope, prefix, base = _scope_for(root)

    _purge(scope)

    file_id = _new_id()
    stored_name = f"{file_id}_{_sanitize_filename(filename)}"

    store = _object_store()
    key = _key(prefix, "databases", stored_name)
    store.put(key, content)

    inspected = inspect_upload(store.materialize(key))

    _save_import(
        file_id,
        scope,
        prefix,
        "database",
        filename,
        stored_name,
        "databases",
        len(content),
        inspected,
        {},
    )

    return {
        "id": file_id,
        "kind": "database",
        "filename": filename,
        "path": f"databases/{stored_name}",
        **inspected,
    }


# ============================================================
# SQL SCRIPTS (SANDBOX ONLY, NEVER store.db)
# ============================================================

def _split_statements(text):
    """Split SQL text on top-level semicolons, ignoring strings/-- lines."""

    statements = []
    current = []
    in_single = False
    in_double = False

    for line in text.splitlines():

        stripped = line.strip()

        if not in_single and not in_double and stripped.startswith("--"):
            continue

        for char in line:

            if char == "'" and not in_double:
                in_single = not in_single
            elif char == '"' and not in_single:
                in_double = not in_double

            current.append(char)

            if char == ";" and not in_single and not in_double:
                statement = "".join(current).strip()
                if statement:
                    statements.append(statement)
                current = []

    remainder = "".join(current).strip()
    if remainder:
        statements.append(remainder)

    return statements


def import_sql_script(filename, content, root=None):
    """Execute an uploaded SQL script into a fresh sandbox database."""

    _ensure_upload_size(content)

    scope, prefix, base = _scope_for(root)

    _purge(scope)

    try:
        text = content.decode("utf-8")
    except UnicodeDecodeError:
        text = content.decode("utf-8", errors="replace")

    statements = _split_statements(text)

    if not statements:
        raise ValidationError("The SQL script contains no statements.")

    if len(statements) > config.MAX_SQL_SCRIPT_STATEMENTS:
        raise ValidationError(
            f"SQL script has too many statements "
            f"(max {config.MAX_SQL_SCRIPT_STATEMENTS})."
        )

    file_id = _new_id()
    stored_name = f"{file_id}_sandbox.db"

    # The database is built in a throwaway directory and only the
    # finished bytes are handed to storage, so a failure part-way
    # through leaves nothing behind in the namespace.
    with tempfile.TemporaryDirectory(
        prefix="tiix-import-"
    ) as workspace:

        target = os.path.join(workspace, "sandbox.db")

        connection = sqlite3.connect(target)
        connection.execute("PRAGMA foreign_keys = ON")

        try:

            for statement in statements:

                if not _STATEMENT_START_RE.match(statement):
                    raise ValidationError(
                        "Only CREATE TABLE/INDEX/VIEW and INSERT INTO "
                        "statements are allowed in imported SQL scripts. "
                        f"Rejected: {statement[:80]}"
                    )

                lower = statement.lower()
                if any(
                    word in lower
                    for word in [
                        "attach",
                        "detach",
                        "vacuum",
                        "load_extension",
                    ]
                ):
                    raise ValidationError(
                        "Import scripts cannot ATTACH/VACUUM or load "
                        "extensions."
                    )

                connection.execute(statement)

            connection.commit()

        except Exception:

            connection.rollback()
            connection.close()
            raise

        else:

            connection.close()

        inspected = inspect_upload(target)
        data = _read_object_bytes(target)

    store = _object_store()
    key = _key(prefix, "sandboxes", stored_name)
    store.put(key, data)

    _save_import(
        file_id,
        scope,
        prefix,
        "sql_script",
        filename,
        stored_name,
        "sandboxes",
        len(data),
        inspected,
        {"statement_count": len(statements)},
    )

    return {
        "id": file_id,
        "kind": "sql_script",
        "filename": filename,
        "path": f"sandboxes/{stored_name}",
        "statement_count": len(statements),
        **inspected,
    }


# ============================================================
# EXCEL FILES
# ============================================================

def _sanitize_table_name(sheet_name, index):

    value = re.sub(r"[^A-Za-z0-9_]", "_", sheet_name).strip("_")

    if not value:
        value = f"sheet_{index}"

    if not re.match(r"^[A-Za-z]", value):
        value = f"s_{value}"

    return value


def _unique_header(column_index, name):

    if not _TABLE_NAME_RE.match(name):
        name = f"column_{column_index}"

    return name


def _infer_column_type(values):

    non_null = [v for v in values if v is not None]

    if non_null and all(
        isinstance(v, int) and not isinstance(v, bool) for v in non_null
    ):
        return "INTEGER"

    if non_null and all(
        isinstance(v, (int, float)) and not isinstance(v, bool)
        for v in non_null
    ):
        return "REAL"

    return "TEXT"


def _normalize_cell(value):

    if value is None:
        return None

    if isinstance(value, bool):
        return 1 if value else 0

    if isinstance(value, (int, float)):
        return value

    return str(value)


def import_excel(filename, content, root=None):
    """Convert an .xlsx workbook into a sandbox SQLite database."""

    _ensure_upload_size(content)

    scope, prefix, base = _scope_for(root)

    _purge(scope)

    from openpyxl import load_workbook

    workbook = load_workbook(
        io.BytesIO(content),
        read_only=True,
        data_only=True,
    )

    file_id = _new_id()
    stored_name = f"{file_id}_excel.db"

    with tempfile.TemporaryDirectory(
        prefix="tiix-import-"
    ) as workspace:

        target = os.path.join(workspace, "excel.db")

        connection = sqlite3.connect(target)

        try:

            sheet_count = 0
            row_count = 0
            truncated = False

            for index, sheet in enumerate(
                workbook.worksheets, start=1
            ):

                rows = sheet.iter_rows(values_only=True)

                try:
                    header_row = next(rows)
                except StopIteration:
                    continue

                header = []
                seen = set()

                for column_index, raw in enumerate(
                    header_row, start=1
                ):
                    name = _unique_header(
                        column_index,
                        re.sub(
                            r"\s+", "_", str(raw or "")
                        ).strip("_"),
                    )
                    while name in seen:
                        name = f"{name}_2"
                    seen.add(name)
                    header.append(name)

                if not header:
                    continue

                data_rows = []

                for values in rows:

                    if len(data_rows) >= config.MAX_IMPORT_ROWS:
                        # The cap is reached and there is at least one
                        # more row to read, so data is being dropped.
                        # Recorded rather than left silent.
                        truncated = True
                        break

                    data_rows.append(
                        [
                            _normalize_cell(v)
                            for v in values[: len(header)]
                        ]
                        + [None] * max(0, len(header) - len(values))
                    )

                if not data_rows:
                    continue

                columns = []

                for column_index, name in enumerate(header):
                    column_values = [
                        row[column_index] for row in data_rows
                    ]
                    columns.append(
                        f"{name} "
                        f"{_infer_column_type(column_values)}"
                    )

                table_name = _sanitize_table_name(sheet.title, index)

                connection.execute(
                    f"CREATE TABLE {table_name} "
                    f"({', '.join(columns)})"
                )

                placeholders = ", ".join(["?"] * len(header))
                connection.executemany(
                    f"INSERT INTO {table_name} "
                    f"({', '.join(header)}) VALUES ({placeholders})",
                    [tuple(row) for row in data_rows],
                )

                sheet_count += 1
                row_count += len(data_rows)

            connection.commit()

        except Exception:

            connection.rollback()
            connection.close()
            workbook.close()
            raise

        else:

            connection.close()
            workbook.close()

        if sheet_count == 0:
            raise ValidationError(
                "The Excel file contains no importable sheets."
            )

        inspected = inspect_upload(target)
        data = _read_object_bytes(target)

    store = _object_store()
    key = _key(prefix, "sandboxes", stored_name)
    store.put(key, data)

    _save_import(
        file_id,
        scope,
        prefix,
        "excel",
        filename,
        stored_name,
        "sandboxes",
        len(data),
        inspected,
        {
            "sheet_count": sheet_count,
            "row_count": row_count,
            "truncated": truncated,
        },
    )

    return {
        "id": file_id,
        "kind": "excel",
        "filename": filename,
        "path": f"sandboxes/{stored_name}",
        "sheet_count": sheet_count,
        "row_count": row_count,
        "truncated": truncated,
        **inspected,
    }


# ============================================================
# CSV FILES
# ============================================================
# A CSV is decoded, sniffed for a delimiter, and materialised as one
# sandbox SQLite table. The result reuses the same column-name
# sanitising and type inference as the Excel path, so the two import
# kinds produce the same shape of database and the same schema text.

# Only a plainly numeric token is coerced. A leading zero (an id or a
# phone number) and an underscore-separated token are left as text, so
# formatting is never silently changed by the import.
_CSV_INT_RE = re.compile(r"^-?(0|[1-9][0-9]*)$")
_CSV_FLOAT_RE = re.compile(r"^-?(0|[1-9][0-9]*)\.[0-9]+$")

# Enough of the file for csv.Sniffer to detect a delimiter without
# reading the whole upload twice.
_CSV_SAMPLE_CHARS = 4096


def _decode_text(content):
    """Decode an uploaded text file, preferring UTF-8.

    utf-8-sig strips a BOM when Excel wrote one. cp1252 and latin-1 are
    tried so a Windows-exported file is still readable; latin-1 cannot
    fail, so a final fallback is not reachable.
    """

    for encoding in ("utf-8-sig", "utf-8", "cp1252", "latin-1"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue

    return content.decode("utf-8", errors="replace")


def _read_csv_rows(text):
    """Parse CSV text into a list of non-empty rows."""

    try:
        dialect = csv.Sniffer().sniff(
            text[:_CSV_SAMPLE_CHARS],
            delimiters=",;\t|",
        )
    except csv.Error:
        dialect = csv.excel

    reader = csv.reader(io.StringIO(text), dialect)

    return [row for row in reader if row]


def _clean_headers(raw_header):
    """Sanitise and uniquify a CSV header row.

    Each name is trimmed, whitespace collapsed to underscores, forced
    through the shared identifier rule, and made unique against the
    names already used in this table.
    """

    header = []
    seen = set()

    for index, raw in enumerate(raw_header, start=1):

        name = _unique_header(
            index,
            re.sub(r"\s+", "_", str(raw or "")).strip("_"),
        )

        base = name
        suffix = 2

        while name in seen:
            name = f"{base}_{suffix}"
            suffix += 1

        seen.add(name)
        header.append(name)

    return header


def _coerce_csv_value(value):

    if value is None:
        return None

    text = str(value).strip()

    if text == "":
        return None

    if _CSV_INT_RE.match(text):
        return int(text)

    if _CSV_FLOAT_RE.match(text):
        return float(text)

    return text


def import_csv(filename, content, root=None):
    """Convert an uploaded .csv into a sandbox SQLite database."""

    _ensure_upload_size(content)

    scope, prefix, base = _scope_for(root)

    _purge(scope)

    rows = _read_csv_rows(_decode_text(content))

    if not rows:
        raise ValidationError("The CSV file is empty.")

    header = _clean_headers(rows[0])

    if not header:
        raise ValidationError("The CSV file has no columns.")

    data_rows = rows[1:]

    if not data_rows:
        raise ValidationError("The CSV file has no data rows.")

    truncated = len(data_rows) > config.MAX_IMPORT_ROWS

    if truncated:
        data_rows = data_rows[:config.MAX_IMPORT_ROWS]

    normalized = []

    for values in data_rows:

        # A short row is padded and a long row is trimmed to the
        # header width, so an inconsistent file still loads and no
        # column silently shifts into another.
        row = [
            _coerce_csv_value(value)
            for value in values[:len(header)]
        ]
        row += [None] * (len(header) - len(row))
        normalized.append(row)

    columns = [
        f"{name} "
        f"{_infer_column_type([row[index] for row in normalized])}"
        for index, name in enumerate(header)
    ]

    table_name = _sanitize_table_name(
        os.path.splitext(os.path.basename(filename))[0],
        1,
    )

    file_id = _new_id()
    stored_name = f"{file_id}_csv.db"

    with tempfile.TemporaryDirectory(
        prefix="tiix-import-"
    ) as workspace:

        target = os.path.join(workspace, "csv.db")

        connection = sqlite3.connect(target)

        try:

            connection.execute(
                f"CREATE TABLE {table_name} "
                f"({', '.join(columns)})"
            )

            placeholders = ", ".join(["?"] * len(header))

            connection.executemany(
                f"INSERT INTO {table_name} "
                f"({', '.join(header)}) VALUES ({placeholders})",
                [tuple(row) for row in normalized],
            )

            connection.commit()

        except Exception:

            connection.rollback()
            connection.close()
            raise

        else:

            connection.close()

        inspected = inspect_upload(target)
        data = _read_object_bytes(target)

    store = _object_store()
    key = _key(prefix, "sandboxes", stored_name)
    store.put(key, data)

    _save_import(
        file_id,
        scope,
        prefix,
        "csv",
        filename,
        stored_name,
        "sandboxes",
        len(data),
        inspected,
        {
            "row_count": len(normalized),
            "truncated": truncated,
        },
    )

    return {
        "id": file_id,
        "kind": "csv",
        "filename": filename,
        "path": f"sandboxes/{stored_name}",
        "row_count": len(normalized),
        "truncated": truncated,
        **inspected,
    }


# ============================================================
# ROUTER
# ============================================================

def import_file(filename, content, root=None):
    """Dispatch an upload by extension. Never touches store.db.

    ``root`` selects the namespace. It defaults to the shared upload
    root (the owner path); a visitor passes their own visitor root,
    which confines the resulting file to their namespace.
    """

    extension = os.path.splitext(filename)[1].lower()

    if extension in config.ALLOWED_DATABASE_EXTENSIONS:
        return import_database(filename, content, root=root)

    if extension in config.ALLOWED_SQL_EXTENSIONS:
        return import_sql_script(filename, content, root=root)

    if extension in config.ALLOWED_EXCEL_EXTENSIONS:
        return import_excel(filename, content, root=root)

    if extension in config.ALLOWED_CSV_EXTENSIONS:
        return import_csv(filename, content, root=root)

    raise ValidationError(
        "Unsupported file type. "
        "Use .db/.sqlite/.sqlite3, .sql, .xlsx, or .csv."
    )


def resolve_import_path(import_id, root=None):
    """Local path of a stored import, confined to the given namespace.

    Raises ValidationError for a malformed id and ImportNotFound when a
    well-formed id matches no recorded import in this namespace or its
    object is missing. It never falls back to the caller's default: an
    explicit import_id either resolves to a stored file or fails.

    The local path is produced by ObjectStore.materialize, which for a
    remote backend downloads the object to a local cache. That is the
    point of the indirection: the SQLite read path needs a real file,
    and where the bytes actually live is the backend's concern.
    """

    if not import_id or not _IMPORT_ID_RE.match(import_id):
        raise ValidationError("Invalid import id.")

    scope, prefix, base = _scope_for(root)

    record = _metadata_store().get(scope, import_id)

    if not record or record.get("scope") != scope:
        raise ImportNotFound("Import not found.")

    store = _object_store()
    key = record.get("object_key")

    if not key or not store.exists(key):
        raise ImportNotFound("Import not found.")

    return store.materialize(key)


def delete_import(import_id, root=None):
    """Remove an import's object and metadata record from a namespace."""

    if not import_id or not _IMPORT_ID_RE.match(import_id):
        return False

    scope, prefix, base = _scope_for(root)

    metadata = _metadata_store()
    record = metadata.get(scope, import_id)

    if not record or record.get("scope") != scope:
        return False

    key = record.get("object_key")

    if key:
        try:
            _object_store().delete(key)
        except storage.StorageError:
            pass

    metadata.delete(scope, import_id)

    return True


def get_import(import_id, root=None):
    """Look up an import by id, inspecting it on demand."""

    if not import_id or not _IMPORT_ID_RE.match(import_id):
        return None

    scope, prefix, base = _scope_for(root)

    record = _metadata_store().get(scope, import_id)

    if not record or record.get("scope") != scope:
        return None

    store = _object_store()
    key = record.get("object_key")

    if not key or not store.exists(key):
        return None

    return {
        "id": import_id,
        "kind": record.get("kind"),
        "filename": record.get("stored_name"),
        "path": record.get("path"),
        **inspect_upload(store.materialize(key)),
    }


def list_imports(root=None):
    """List stored imports without re-inspecting each file."""

    scope, prefix, base = _scope_for(root)

    records = [
        record
        for record in _metadata_store().list(scope)
        if record.get("scope") == scope
    ]

    records.sort(
        key=lambda record: (
            record.get("created_at", 0),
            record.get("id", ""),
        )
    )

    return [
        {
            "id": record.get("id"),
            "kind": record.get("kind"),
            "filename": record.get("stored_name"),
            "path": record.get("path"),
        }
        for record in records
    ]


# ============================================================
# EXPIRY
# ============================================================
# A visitor identity expires with its cookie; the files it uploaded
# should not outlive it. There is no scheduler here, so expiry is
# swept on the next import in the same namespace: the operation that
# adds a record is also the one that removes the stale ones, which
# keeps the store bounded without a background job the platform may
# not run. Zero disables the sweep.

def _purge(scope, now=None):
    """Delete records older than UPLOAD_TTL_SECONDS in one namespace."""

    ttl = int(getattr(config, "UPLOAD_TTL_SECONDS", 0) or 0)

    if ttl <= 0:
        return 0

    current = float(now) if now is not None else time.time()

    metadata = _metadata_store()
    objects = _object_store()

    removed = 0

    for record in metadata.list(scope):

        if record.get("scope") != scope:
            continue

        created = record.get("created_at")

        if not isinstance(created, (int, float)):
            continue

        if current - created <= ttl:
            continue

        key = record.get("object_key")

        if key:
            try:
                objects.delete(key)
            except storage.StorageError:
                pass

        metadata.delete(scope, record.get("id"))
        removed += 1

    return removed


def purge_expired_imports(root=None, now=None):
    """Sweep expired imports in the given namespace (default: shared)."""

    scope, prefix, base = _scope_for(root)

    return _purge(scope, now=now)


# ============================================================
# VISITOR NAMESPACE
# ============================================================
# Thin wrappers that pin every call to a single visitor's root. They
# exist so a route can never forget the root argument and fall back to
# the shared namespace by accident: the only way to reach a visitor's
# files is through a function whose name says "visitor".

def import_file_for_visitor(visitor_id, filename, content):
    return import_file(
        filename, content, root=_visitor_root(visitor_id)
    )


def list_visitor_imports(visitor_id):
    return list_imports(root=_visitor_root(visitor_id))


def resolve_visitor_import_path(visitor_id, import_id):
    return resolve_import_path(
        import_id, root=_visitor_root(visitor_id)
    )


def delete_visitor_import(visitor_id, import_id):
    return delete_import(
        import_id, root=_visitor_root(visitor_id)
    )
